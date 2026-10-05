import unittest
from pathlib import Path
from tempfile import TemporaryDirectory


class TestBusinessTypes(unittest.TestCase):
    def test_detect_kinds(self) -> None:
        from docs.business_types import detect_business_kind

        self.assertEqual(
            detect_business_kind("Составь договор оказания услуг с ООО Ромашка"),
            "contract",
        )
        self.assertEqual(
            detect_business_kind("Нужен акт выполненных работ за март"),
            "act",
        )
        self.assertEqual(
            detect_business_kind("Выстави счёт на 50000 за разработку"),
            "invoice",
        )
        self.assertEqual(
            detect_business_kind("Напиши претензию поставщику о просрочке"),
            "claim",
        )
        self.assertEqual(
            detect_business_kind(
                "Напиши деловое письмо клиенту с просьбой уточнить срок"
            ),
            "letter",
        )

    def test_referat_not_business(self) -> None:
        from docs.business_types import detect_business_kind

        self.assertIsNone(
            detect_business_kind("Напиши реферат на тему ИИ в медицине")
        )

    def test_missing_contract_fields(self) -> None:
        from docs.business_types import missing_fields

        miss = missing_fields(
            "contract", {"party_a": "ИП Иванов", "subject": "сайт"}
        )
        self.assertIn("party_b", miss)
        self.assertIn("price", miss)
        self.assertIn("term", miss)

    def test_claim_default_deadline(self) -> None:
        from docs.business_types import apply_defaults

        out = apply_defaults(
            "claim",
            {
                "addressee": "ООО X",
                "violation_summary": "не оплатил",
                "demand": "оплатить",
            },
        )
        self.assertIn("10 календарных", out.get("reply_deadline", ""))

    def test_multi_kind_picks_first(self) -> None:
        from docs.business_types import detect_business_kind

        kind = detect_business_kind(
            "Сделай договор и потом счёт и акт с ООО Ромашка"
        )
        self.assertEqual(kind, "contract")

    def test_waive_bank_phrases(self) -> None:
        from docs.business_types import wants_waive_bank

        self.assertTrue(wants_waive_bank("реквизитов нет, сделай без них"))

    def test_vague_reply(self) -> None:
        from docs.business_types import is_vague_user_reply

        self.assertTrue(is_vague_user_reply("не знаю"))
        self.assertFalse(is_vague_user_reply("ООО Ромашка, 80 000 руб"))

    def test_kind_from_short_hint(self) -> None:
        from docs.business_types import kind_from_short_hint

        self.assertEqual(kind_from_short_hint("счёт"), "invoice")
        self.assertEqual(kind_from_short_hint("договор"), "contract")


class TestBusinessRender(unittest.TestCase):
    def test_invoice_has_table_and_no_invented_inn(self) -> None:
        from docx import Document
        from edit.business_docx import render_business_docx

        fields = {
            "seller": "ИП Иванов",
            "buyer": "ООО Ромашка",
            "line_items": [
                {"name": "Разработка сайта", "qty": 1, "price": 50000}
            ],
            "bank_details": {
                "bank": "Т-Банк",
                "bik": "044525974",
                "rs": "40802810900000000000",
            },
        }
        with TemporaryDirectory() as td:
            dest = Path(td) / "invoice.docx"
            render_business_docx("invoice", fields, dest)
            doc = Document(str(dest))
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("Счёт", text)
            self.assertTrue(doc.tables)
            self.assertNotIn("выдуман", text.lower())

    def test_contract_contains_parties_and_price(self) -> None:
        from docx import Document
        from edit.business_docx import render_business_docx

        fields = {
            "party_a": "ИП Иванов",
            "party_b": "ООО Ромашка",
            "subject": "разработка сайта",
            "price": "80000 руб.",
            "term": "30 календарных дней",
        }
        with TemporaryDirectory() as td:
            dest = Path(td) / "contract.docx"
            render_business_docx("contract", fields, dest)
            doc = Document(str(dest))
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("ИП Иванов", text)
            self.assertIn("ООО Ромашка", text)
            self.assertIn("80000", text)

    def test_act_claim_letter_smoke(self) -> None:
        from docx import Document
        from edit.business_docx import render_business_docx

        cases = [
            (
                "act",
                {
                    "party_a": "ИП А",
                    "party_b": "ООО Б",
                    "services": ["Консультация"],
                    "amount": "10000",
                    "period": "март 2026",
                },
                "Акт",
            ),
            (
                "claim",
                {
                    "addressee": "ООО Должник",
                    "violation_summary": "не оплатил счёт",
                    "demand": "оплатить долг",
                },
                "Претензия",
            ),
            (
                "letter",
                {
                    "addressee": "ООО Клиент",
                    "body_purpose": "Просим уточнить срок поставки.",
                    "sender_sign": "ИП Иванов",
                },
                "письмо",
            ),
        ]
        with TemporaryDirectory() as td:
            for kind, fields, needle in cases:
                dest = Path(td) / f"{kind}.docx"
                render_business_docx(kind, fields, dest)
                doc = Document(str(dest))
                text = "\n".join(p.text for p in doc.paragraphs).lower()
                self.assertIn(needle.lower(), text)


class TestBusinessExtract(unittest.TestCase):
    def test_sanitize_drops_invented_inn(self) -> None:
        from llm.business_extract import sanitize_business_fields

        prompt = "Счёт от ИП Иванов для ООО Ромашка, 10000 руб за консультацию"
        raw = {
            "seller": "ИП Иванов",
            "buyer": "ООО Ромашка",
            "inn_seller": "7707083893",
            "line_items": [{"name": "консультация", "qty": 1, "price": 10000}],
        }
        clean = sanitize_business_fields(prompt, "invoice", raw)
        self.assertFalse(clean.get("inn_seller") or clean.get("inn"))

    def test_merge_answers_into_fields(self) -> None:
        from llm.business_extract import merge_extracted_fields

        base = {"seller": "ИП Иванов"}
        merged = merge_extracted_fields(
            base, {"buyer": "ООО Ромашка", "amount": "50000"}
        )
        self.assertEqual(merged["buyer"], "ООО Ромашка")


class TestBusinessDraftSession(unittest.TestCase):
    def test_clear_dialog_modes_clears_draft(self) -> None:
        from bot.session import BusinessDraft, UserSession

        s = UserSession(user_id=1)
        s.business_draft = BusinessDraft(kind="invoice", base_prompt="счёт")
        s.clear_dialog_modes()
        self.assertIsNone(s.business_draft)

    def test_payload_roundtrip(self) -> None:
        from bot.session import BusinessDraft, UserSession
        from bot.session_store import apply_payload, session_to_payload

        s = UserSession(user_id=1)
        s.business_draft = BusinessDraft(
            kind="act",
            fields={"amount": "1"},
            missing=["period"],
            base_prompt="x",
        )
        payload = session_to_payload(s)
        s2 = UserSession(user_id=1)
        apply_payload(s2, payload)
        assert s2.business_draft is not None
        self.assertEqual(s2.business_draft.kind, "act")
        self.assertEqual(s2.business_draft.missing, ["period"])


class TestBusinessIntent(unittest.TestCase):
    def test_compose_contract_is_business_write(self) -> None:
        from llm.client import classify_document_intent

        d = classify_document_intent(
            "Составь договор оказания услуг с ООО Ромашка на 80 тысяч, срок месяц",
            has_files=False,
        )
        self.assertEqual(d["intent"], "business_write")
        self.assertEqual(d["doc_kind"], "contract")

    def test_referat_still_write_text(self) -> None:
        from llm.client import classify_document_intent

        d = classify_document_intent("Напиши реферат на тему ИИ", has_files=False)
        self.assertEqual(d["intent"], "write_text")


if __name__ == "__main__":
    unittest.main()
