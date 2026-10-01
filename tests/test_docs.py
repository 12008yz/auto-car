"""
Автотесты без Telegram и без вызовов LLM.
Запуск из корня проекта:
  .venv\\Scripts\\python.exe -m unittest discover -s tests -v
"""

from __future__ import annotations

import ast
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


class TestSyntax(unittest.TestCase):
    def test_core_modules_parse(self) -> None:
        files = [
            "bot/handlers.py",
            "bot/ui.py",
            "bot/session.py",
            "llm/client.py",
            "edit/docx_patch.py",
            "billing/service.py",
            "config.py",
        ]
        for rel in files:
            path = ROOT / rel
            with self.subTest(file=rel):
                ast.parse(path.read_text(encoding="utf-8"))


class TestIntentRouter(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        from llm.client import (
            classify_document_intent,
            looks_like_apply_pending,
            looks_like_card,
            looks_like_edit,
            looks_like_fill_data,
            looks_like_reverse_words,
            looks_like_tone,
        )

        cls.classify = staticmethod(classify_document_intent)
        cls.looks_like_card = staticmethod(looks_like_card)
        cls.looks_like_edit = staticmethod(looks_like_edit)
        cls.looks_like_fill_data = staticmethod(looks_like_fill_data)
        cls.looks_like_tone = staticmethod(looks_like_tone)
        cls.looks_like_reverse_words = staticmethod(looks_like_reverse_words)
        cls.looks_like_apply_pending = staticmethod(looks_like_apply_pending)

    def _intent(self, text: str, *, has_files: bool = False) -> str:
        return str(self.classify(text, has_files=has_files)["intent"])

    def _mode(self, text: str, *, has_files: bool = False) -> str:
        return str(self.classify(text, has_files=has_files).get("mode") or "")

    def test_fill_declaration_is_edit_not_card(self) -> None:
        text = (
            "Нужно добавить данные: Чикасов Денис Владимирович "
            "27.11.1999 года Богородицк"
        )
        self.assertTrue(self.looks_like_fill_data(text) or self.looks_like_edit(text))
        self.assertFalse(self.looks_like_card(text))
        self.assertEqual(self._intent(text, has_files=True), "edit")
        self.assertEqual(self._mode(text, has_files=True), "fill")

    def test_fio_followup_is_edit(self) -> None:
        text = "Сделай фио Чикасов Денис Владимирович"
        self.assertEqual(self._intent(text, has_files=True), "edit")

    def test_add_these_data_is_edit(self) -> None:
        self.assertEqual(
            self._intent("Нужно добавить эти данные", has_files=True), "edit"
        )

    def test_contract_question_is_ask_not_extract(self) -> None:
        self.assertEqual(
            self._intent("Что в договоре про срок оплаты?", has_files=True), "ask"
        )

    def test_risks_extract_compare_format_check(self) -> None:
        cases = [
            ("Проверь договор на риски", True, "risks"),
            ("Вытащи стороны, даты и суммы", True, "extract"),
            ("Сравни эти два файла", True, "compare"),
            ("Оформи по образцу", True, "format"),
            ("Что не заполнено в бланке?", True, "check"),
        ]
        for text, has_files, expected in cases:
            with self.subTest(text=text):
                self.assertEqual(self._intent(text, has_files=has_files), expected)

    def test_tone_mode(self) -> None:
        text = "Сделай текст официальнее"
        self.assertTrue(self.looks_like_tone(text))
        self.assertEqual(self._intent(text, has_files=True), "edit")
        self.assertEqual(self._mode(text, has_files=True), "tone")

    def test_write_and_card(self) -> None:
        self.assertEqual(
            self._intent("Напиши реферат на тему ИИ", has_files=False), "write_text"
        )
        self.assertEqual(
            self._intent("Напиши декларацию для ИП", has_files=False), "clarify"
        )
        self.assertEqual(
            self._intent("Сделай титульный лист декларации ИП", has_files=False),
            "write_form",
        )
        self.assertEqual(
            self._intent("Составь претензию поставщику о просрочке", has_files=False),
            "write_text",
        )
        self.assertEqual(
            self._intent("Нужно продать кружку керамическую 300 мл", has_files=False),
            "card",
        )

    def test_replace_is_edit(self) -> None:
        self.assertEqual(
            self._intent("Замени Иванова на Петрова", has_files=True), "edit"
        )

    def test_soft_need_verify_with_files_is_ask_not_clarify(self) -> None:
        phrases = [
            "Мне нужно что бы ты проверил информацию",
            "Нужно чтобы ты проверил информацию о человеке",
            "Проверь информацию в файле",
            "Посмотри что внутри",
        ]
        for text in phrases:
            with self.subTest(text=text):
                d = self.classify(text, has_files=True)
                self.assertEqual(d["intent"], "ask", msg=d)
                self.assertNotEqual(d.get("family"), "write")

    def test_soft_need_without_strong_create_not_write_clarify(self) -> None:
        text = "Нужно про декларацию"
        d = self.classify(text, has_files=True)
        self.assertNotEqual(d["intent"], "clarify")
        self.assertNotEqual(d.get("family"), "write")

    def test_verify_without_files_asks_for_file(self) -> None:
        d = self.classify("Проверь информацию", has_files=False)
        self.assertEqual(d["intent"], "clarify")
        self.assertEqual(d.get("family"), "ask")
        q = str(d.get("question") or "").lower()
        self.assertTrue("документ" in q or "файл" in q)

    def test_strong_write_with_files_still_write(self) -> None:
        self.assertEqual(
            self._intent("Напиши реферат на тему что такое ИИ", has_files=True),
            "write_text",
        )
        self.assertEqual(
            self._intent("Создай новый документ — претензию", has_files=True),
            "write_text",
        )

    def test_guard_metadata_present_on_overrides(self) -> None:
        # Сырой роутер даёт write_form/clarify; guard обязан переписать при файлах
        d = self.classify("Нужна декларация", has_files=True)
        self.assertEqual(d["intent"], "ask")
        self.assertTrue(d.get("guard"))
        self.assertNotEqual(d.get("family"), "write")

    def test_reverse_words_is_edit_not_fill(self) -> None:
        phrases = [
            "Перепиши слова задом наперёд",
            "Нужно перевернуть все слова в файле и скинуть мне его",
            "файл мне дай с перевёрнутыми словами",
        ]
        for text in phrases:
            with self.subTest(text=text):
                self.assertTrue(self.looks_like_reverse_words(text), msg=text)
                d = self.classify(text, has_files=True)
                self.assertEqual(d["intent"], "edit", msg=d)
                self.assertNotEqual(d.get("mode"), "fill", msg=d)

    def test_apply_pending_phrases(self) -> None:
        for text in ("Давай", "Ну ты мне файл дай", "скинь файл", "примени"):
            with self.subTest(text=text):
                self.assertTrue(self.looks_like_apply_pending(text), msg=text)
        self.assertFalse(self.looks_like_apply_pending("год должен быть 9991"))
        self.assertFalse(self.looks_like_apply_pending("Замени Иванова на Петрова"))


class TestPersonFacts(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        from llm.client import _extract_person_facts, _sanitize_fill_patches

        cls.extract = staticmethod(_extract_person_facts)
        cls.sanitize = staticmethod(_sanitize_fill_patches)

    def test_extract_chikasov_phrase(self) -> None:
        text = (
            "Нужно добавить данные: Чикасов Денис Владимирович "
            "27.11.1999 года Богородицк"
        )
        facts = self.extract(text)
        self.assertTrue(facts.get("fio", "").startswith("Чикасов"))
        self.assertEqual(facts.get("birth_date"), "27.11.1999")
        self.assertEqual(facts.get("birth_year"), "1999")
        self.assertEqual(facts.get("city"), "Богородицк")
        self.assertEqual(facts.get("date_role"), "birth")

    def test_sanitize_drops_birth_year_in_report_year_field(self) -> None:
        facts = {
            "fio": "Чикасов Денис Владимирович",
            "birth_date": "27.11.1999",
            "birth_year": "1999",
            "city": "Богородицк",
            "date_role": "birth",
        }
        bad = [
            {"find": "Год: [ГГГГ]", "replace": "1999"},
            {"find": "[ФИО]", "replace": facts["fio"]},
        ]
        hits = (
            "ФИО: [ФИО]\nГод: [ГГГГ]\n"
            "Дата рождения: [ДД.ММ.ГГГГ]\nГород: [город]\n"
        )
        good = self.sanitize(bad, facts, hits)
        self.assertFalse(any(p["replace"] == "1999" for p in good))
        self.assertTrue(any(facts["fio"] in p["replace"] for p in good))


class TestDocSession(unittest.TestCase):
    def test_fill_instruction_merges_facts(self) -> None:
        from bot.session import UserSession

        session = UserSession(user_id=1)
        session.remember_doc_task(
            "fill",
            "Чикасов Денис Владимирович 27.11.1999 Богородицк",
            accumulate_facts=True,
        )
        merged = session.fill_instruction("Нужно добавить эти данные")
        self.assertIn("Чикасов", merged)
        self.assertIn("эти данные", merged.lower())

    def test_reset_clears_doc_task(self) -> None:
        from bot.session import UserSession

        session = UserSession(user_id=2)
        session.remember_doc_task("fill", "тест", accumulate_facts=True)
        session.reset_memory()
        self.assertIsNone(session.doc_task)


class TestDocxWriter(unittest.TestCase):
    def test_write_form_and_text_layouts(self) -> None:
        from edit.docx_patch import write_structured_docx
        from docx import Document

        sections = [
            {
                "heading": "Данные",
                "paragraphs": ["ФИО: [ФИО]", "Дата рождения: [ДД.ММ.ГГГГ]"],
            }
        ]
        with tempfile.TemporaryDirectory() as tmp:
            form_path = Path(tmp) / "form.docx"
            text_path = Path(tmp) / "text.docx"
            write_structured_docx(
                form_path,
                title="Титульный",
                sections=sections,
                doc_type="титульный",
                layout="form",
            )
            write_structured_docx(
                text_path,
                title="Реферат",
                sections=[{"heading": "Введение", "paragraphs": ["Текст абзаца."]}],
                doc_type="реферат",
                layout="text",
            )
            self.assertTrue(form_path.exists())
            self.assertTrue(text_path.exists())
            form_doc = Document(str(form_path))
            text_doc = Document(str(text_path))
            self.assertGreaterEqual(len([p for p in form_doc.paragraphs if p.text.strip()]), 2)
            self.assertGreaterEqual(len([p for p in text_doc.paragraphs if p.text.strip()]), 2)


class TestReverseWords(unittest.TestCase):
    def test_reverse_words_in_text(self) -> None:
        from edit.docx_patch import reverse_words_in_text

        self.assertEqual(
            reverse_words_in_text("ФИО: Чикасов Денис Год: 1999"),
            "ОИФ: восакиЧ синеД доГ: 9991",
        )

    def test_reverse_words_docx(self) -> None:
        from docx import Document
        from edit.docx_patch import reverse_words_docx, write_structured_docx

        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "src.docx"
            dest = Path(tmp) / "out.docx"
            write_structured_docx(
                src,
                title="Титульный лист",
                sections=[{"heading": None, "paragraphs": ["Год: 1999"]}],
                doc_type="титульный",
                layout="form",
            )
            result = reverse_words_docx(src, dest)
            self.assertTrue(dest.exists())
            self.assertGreater(int(result["changed"]), 0)
            blob = "\n".join(p.text for p in Document(str(dest)).paragraphs)
            self.assertIn("9991", blob)
            self.assertNotIn("1999", blob)


class TestUiWiring(unittest.TestCase):
    def test_docs_buttons_and_examples(self) -> None:
        from bot import ui

        kb = ui.docs_inline()
        data = {
            btn.callback_data
            for row in kb.inline_keyboard
            for btn in row
            if btn.callback_data
        }
        for need in (
            "menu:risks",
            "menu:extract",
            "menu:compare",
            "menu:gaps",
            "menu:docs_fill_ex",
            "menu:docs_more_ex",
        ):
            self.assertIn(need, data)
        self.assertTrue(ui.docs_fill_example_text())
        self.assertTrue(ui.docs_more_examples_text())
        self.assertIn("бесплатно", ui.help_text().lower())


if __name__ == "__main__":
    unittest.main()
