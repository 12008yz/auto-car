"""Typed .docx renderers for business document drafts."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from docs.business_types import KIND_LABELS_RU, apply_defaults


def suggest_business_filename(kind: str, fields: dict[str, Any] | None = None) -> str:
    data = fields or {}
    label = {
        "contract": "dogovor",
        "act": "akt",
        "invoice": "schet",
        "claim": "pretenziya",
        "letter": "pismo",
    }.get(kind, "document")
    topic = ""
    for key in ("subject", "service_name", "body_purpose", "buyer", "party_b", "addressee"):
        raw = str(data.get(key) or "").strip()
        if raw:
            topic = raw
            break
    stem = _safe_stem(f"{label}_{topic}" if topic else label)
    return f"{stem}.docx"


def render_business_docx(kind: str, fields: dict[str, Any], dest: Path) -> Path:
    kind = (kind or "").strip().lower()
    if kind not in KIND_LABELS_RU:
        raise ValueError(f"Неизвестный тип документа: {kind}")
    data = apply_defaults(kind, fields)
    dest = Path(dest)
    if dest.suffix.lower() != ".docx":
        dest = dest.with_suffix(".docx")
    dest.parent.mkdir(parents=True, exist_ok=True)

    if kind == "invoice":
        _render_invoice(data, dest)
    elif kind == "act":
        _render_act(data, dest)
    elif kind == "contract":
        _render_contract(data, dest)
    elif kind == "claim":
        _render_claim(data, dest)
    else:
        _render_letter(data, dest)
    return dest


def _safe_stem(name: str, fallback: str = "document") -> str:
    cleaned = re.sub(r"[^\w\- а-яА-ЯёЁ]+", "", (name or "").strip(), flags=re.UNICODE)
    cleaned = re.sub(r"\s+", "_", cleaned).strip("._")
    return (cleaned[:60] or fallback)


def _set_normal(doc: Document, size: int = 12) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)


def _p(doc: Document, text: str, *, bold: bool = False, center: bool = False) -> None:
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(6)
    para.paragraph_format.first_line_indent = Pt(0)
    run = para.add_run(str(text or ""))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def _ph(value: Any, label: str) -> str:
    text = str(value or "").strip()
    return text if text else f"[укажите {label}]"


def _line_items(fields: dict[str, Any]) -> list[dict[str, Any]]:
    items = fields.get("line_items")
    if isinstance(items, list) and items:
        out: list[dict[str, Any]] = []
        for row in items:
            if isinstance(row, dict) and str(row.get("name") or "").strip():
                out.append(row)
        if out:
            return out
    name = str(fields.get("service_name") or fields.get("subject") or "Услуга").strip()
    amount = fields.get("amount") or fields.get("price") or fields.get("total") or 0
    return [{"name": name, "qty": 1, "price": amount}]


def _as_float(value: Any) -> float | None:
    try:
        return float(str(value).replace(" ", "").replace(",", "."))
    except (TypeError, ValueError):
        return None


def _money(value: Any) -> str:
    num = _as_float(value)
    if num is None:
        return str(value or "").strip() or "0"
    if abs(num - round(num)) < 1e-9:
        return f"{int(round(num))}"
    return f"{num:.2f}"


def _render_invoice(fields: dict[str, Any], dest: Path) -> None:
    doc = Document()
    _set_normal(doc)
    _p(doc, "Счёт на оплату", bold=True, center=True)
    _p(doc, f"Продавец: {_ph(fields.get('seller'), 'продавца')}")
    _p(doc, f"Покупатель: {_ph(fields.get('buyer'), 'покупателя')}")
    inn = fields.get("inn") or fields.get("inn_seller")
    kpp = fields.get("kpp")
    if inn:
        _p(doc, f"ИНН продавца: {inn}")
    if kpp:
        _p(doc, f"КПП: {kpp}")

    items = _line_items(fields)
    table = doc.add_table(rows=1, cols=4)
    headers = ("№", "Наименование", "Кол-во", "Сумма")
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    total = 0.0
    for idx, row in enumerate(items, start=1):
        cells = table.add_row().cells
        qty = _as_float(row.get("qty")) or 1.0
        price = _as_float(row.get("price")) or 0.0
        line_sum = qty * price
        total += line_sum
        cells[0].text = str(idx)
        cells[1].text = str(row.get("name") or "")
        cells[2].text = _money(qty)
        cells[3].text = _money(line_sum)

    declared = fields.get("total")
    total_text = _money(declared if declared not in (None, "") else total)
    _p(doc, "")
    _p(doc, f"Итого: {total_text} руб.", bold=True)

    vat = str(fields.get("vat_mode") or "unknown")
    if vat == "with_vat":
        _p(doc, "НДС: включён / выделить по факту учёта")
    elif vat == "no_vat":
        _p(doc, "Без НДС")
    if fields.get("payment_due"):
        _p(doc, f"Срок оплаты: {fields.get('payment_due')}")

    bank = fields.get("bank_details") if isinstance(fields.get("bank_details"), dict) else {}
    _p(doc, "Реквизиты для оплаты:", bold=True)
    _p(doc, f"Банк: {_ph(bank.get('bank'), 'банк')}")
    _p(doc, f"БИК: {_ph(bank.get('bik'), 'БИК')}")
    _p(doc, f"Р/с: {_ph(bank.get('rs'), 'расчётный счёт')}")
    doc.save(str(dest))


def _render_act(fields: dict[str, Any], dest: Path) -> None:
    doc = Document()
    _set_normal(doc)
    _p(doc, "Акт оказанных услуг", bold=True, center=True)
    _p(doc, f"Исполнитель: {_ph(fields.get('party_a'), 'исполнителя')}")
    _p(doc, f"Заказчик: {_ph(fields.get('party_b'), 'заказчика')}")
    ref = fields.get("contract_ref")
    _p(doc, f"Основание: {_ph(ref, '№ и дату договора (или «без договора»)')}")
    _p(doc, f"Период: {_ph(fields.get('period'), 'период')}")

    services = fields.get("services")
    if isinstance(services, list):
        rows = [str(x).strip() for x in services if str(x).strip()]
    elif str(services or "").strip():
        rows = [str(services).strip()]
    else:
        rows = []

    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "№"
    table.rows[0].cells[1].text = "Услуга / работа"
    if not rows:
        rows = ["[укажите перечень услуг]"]
    for i, name in enumerate(rows, start=1):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = name

    _p(doc, "")
    _p(doc, f"Сумма: {_ph(fields.get('amount'), 'сумму')} руб.", bold=True)
    phrase = str(fields.get("no_claims_phrase") or "").strip()
    if not phrase:
        phrase = (
            "Услуги оказаны в полном объёме, в установленный срок. "
            "Заказчик претензий по объёму, качеству и срокам не имеет."
        )
    _p(doc, phrase)
    _p(doc, "")
    _p(doc, "Исполнитель _____________ / _____________")
    _p(doc, "Заказчик _____________ / _____________")
    doc.save(str(dest))


def _render_contract(fields: dict[str, Any], dest: Path) -> None:
    doc = Document()
    _set_normal(doc, size=14)
    _p(doc, "Договор оказания услуг", bold=True, center=True)
    _p(doc, "")
    a = _ph(fields.get("party_a"), "исполнителя")
    b = _ph(fields.get("party_b"), "заказчика")
    _p(
        doc,
        f"1. Стороны. Исполнитель: {a}. Заказчик: {b}. "
        "Стороны заключили настоящий договор о нижеследующем.",
    )
    _p(
        doc,
        f"2. Предмет. Исполнитель обязуется оказать услуги: "
        f"{_ph(fields.get('subject'), 'предмет услуг')}, "
        "а Заказчик обязуется принять и оплатить их.",
    )
    _p(
        doc,
        f"3. Срок. Услуги оказываются в срок: {_ph(fields.get('term'), 'срок')}.",
    )
    _p(
        doc,
        f"4. Цена и оплата. Стоимость услуг составляет "
        f"{_ph(fields.get('price'), 'цену')}. "
        f"{str(fields.get('payment_order') or 'Оплата производится на основании счёта и акта.').strip()}",
    )
    acceptance = str(
        fields.get("acceptance_order")
        or "Факт оказания услуг подтверждается актом, подписанным сторонами."
    ).strip()
    _p(doc, f"5. Приёмка. {acceptance}")
    _p(
        doc,
        "6. Прочие условия. Во всём, что не урегулировано договором, стороны "
        "руководствуются законодательством Российской Федерации. "
        "Изменения действительны в письменной форме.",
    )
    _p(doc, "7. Реквизиты и подписи сторон.", bold=True)
    _p(
        doc,
        f"Исполнитель: {a}"
        + (f", ИНН {fields['inn_a']}" if fields.get("inn_a") else "")
        + (f", адрес: {fields['address_a']}" if fields.get("address_a") else ""),
    )
    _p(
        doc,
        f"Заказчик: {b}"
        + (f", ИНН {fields['inn_b']}" if fields.get("inn_b") else "")
        + (f", адрес: {fields['address_b']}" if fields.get("address_b") else ""),
    )
    _p(doc, "")
    _p(doc, "Исполнитель _____________ / _____________")
    _p(doc, "Заказчик _____________ / _____________")
    doc.save(str(dest))


def _render_claim(fields: dict[str, Any], dest: Path) -> None:
    doc = Document()
    _set_normal(doc)
    _p(doc, "Претензия", bold=True, center=True)
    _p(doc, f"Кому: {_ph(fields.get('addressee'), 'адресата')}")
    if fields.get("contract_ref"):
        _p(doc, f"По договору / основанию: {fields.get('contract_ref')}")
    _p(doc, "")
    _p(
        doc,
        "Настоящим сообщаем о следующем нарушении: "
        f"{_ph(fields.get('violation_summary'), 'описание нарушения')}.",
    )
    if fields.get("debt_amount"):
        _p(doc, f"Сумма задолженности / требования: {fields.get('debt_amount')}.")
    _p(
        doc,
        f"Требуем: {_ph(fields.get('demand'), 'требование')}.",
    )
    _p(
        doc,
        "Просим ответить "
        f"{_ph(fields.get('reply_deadline'), 'срок ответа')}. "
        "В случае отказа оставляем за собой право обратиться в суд "
        "и взыскать убытки / неустойку в установленном порядке.",
    )
    if fields.get("attachments_note"):
        _p(doc, f"Приложения: {fields.get('attachments_note')}")
    _p(doc, "")
    _p(doc, "С уважением,")
    _p(doc, _ph(fields.get("sender_sign"), "ФИО / должность"))
    doc.save(str(dest))


def _render_letter(fields: dict[str, Any], dest: Path) -> None:
    doc = Document()
    _set_normal(doc)
    _p(doc, "Деловое письмо", bold=True, center=True)
    _p(doc, f"Кому: {_ph(fields.get('addressee'), 'адресата')}")
    if fields.get("subject_line"):
        _p(doc, f"Тема: {fields.get('subject_line')}")
    _p(doc, "")
    _p(doc, _ph(fields.get("body_purpose"), "суть обращения"))
    if fields.get("reply_deadline"):
        _p(doc, f"Просим ответить: {fields.get('reply_deadline')}.")
    if fields.get("sender_requisites"):
        _p(doc, f"Контакты: {fields.get('sender_requisites')}")
    _p(doc, "")
    _p(doc, "С уважением,")
    _p(doc, _ph(fields.get("sender_sign"), "ФИО / должность"))
    doc.save(str(dest))
