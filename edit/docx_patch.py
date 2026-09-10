from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


def _replace_in_paragraph(paragraph: Paragraph, find: str, replace: str) -> int:
    if find not in paragraph.text:
        return 0
    count = paragraph.text.count(find)
    new_text = paragraph.text.replace(find, replace)
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return count
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""
    return count


def _walk_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def apply_patches(src: Path, dest: Path, patches: list[dict[str, str]]) -> dict:
    doc = Document(str(src))
    applied: list[dict] = []
    missing: list[str] = []
    for patch in patches:
        find = patch["find"]
        replace = patch["replace"]
        total = 0
        for para in _walk_paragraphs(doc):
            total += _replace_in_paragraph(para, find, replace)
        if total:
            applied.append({"find": find, "replace": replace, "count": total})
        else:
            missing.append(find)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return {"applied": applied, "missing": missing, "path": dest}


def write_rewrite_docx(dest: Path, text: str) -> Path:
    doc = Document()
    paragraphs = [p.strip() for p in text.replace("\r\n", "\n").split("\n") if p.strip()]
    if not paragraphs:
        doc.add_paragraph("")
    for p in paragraphs:
        doc.add_paragraph(p)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest
