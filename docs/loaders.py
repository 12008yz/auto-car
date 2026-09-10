from dataclasses import dataclass
from pathlib import Path


@dataclass
class ExtractedBlock:
    text: str
    location: str


def load_document(path: Path) -> list[ExtractedBlock]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _load_text(path)
    if suffix == ".csv":
        return _load_text(path, kind="CSV")
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".xlsx":
        return _load_xlsx(path)
    if suffix == ".pptx":
        return _load_pptx(path)
    raise ValueError(f"Формат не поддерживается: {suffix or path.name}")


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _load_text(path: Path, kind: str = "текст") -> list[ExtractedBlock]:
    text = _read_text_file(path).strip()
    if not text:
        return []
    return [ExtractedBlock(text=text, location=kind)]


def _load_docx(path: Path) -> list[ExtractedBlock]:
    from docx import Document

    doc = Document(str(path))
    blocks: list[ExtractedBlock] = []
    for i, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if text:
            blocks.append(ExtractedBlock(text=text, location=f"абзац {i}"))
    for t_i, table in enumerate(doc.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(
                ExtractedBlock(text="\n".join(rows), location=f"таблица {t_i}")
            )
    return blocks


def _load_pdf(path: Path) -> list[ExtractedBlock]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[ExtractedBlock] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(ExtractedBlock(text=text, location=f"стр. {i}"))
    return blocks


def _load_xlsx(path: Path) -> list[ExtractedBlock]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[ExtractedBlock] = []
    for sheet in wb.worksheets:
        rows: list[str] = []
        for r_i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = ["" if c is None else str(c).strip() for c in row]
            if any(values):
                rows.append(f"{r_i}: " + " | ".join(values))
            if len(rows) >= 400:
                break
        if rows:
            blocks.append(
                ExtractedBlock(
                    text="\n".join(rows),
                    location=f"лист {sheet.title}",
                )
            )
    wb.close()
    return blocks


def _load_pptx(path: Path) -> list[ExtractedBlock]:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[ExtractedBlock] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    parts.append(text)
        if parts:
            blocks.append(ExtractedBlock(text="\n".join(parts), location=f"слайд {i}"))
    return blocks
